import re
import os
import shutil
from pathlib import Path
from datetime import datetime

import pandas as pd
import streamlit as st
from docxtpl import DocxTemplate, RichText
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL

# ============================================================
# KONFIGURASI JALUR BERKAS & TEMPLATE
# ============================================================
BASE_DIR = Path(__file__).resolve().parent

# Berkas Template Word (Telah Dianonimkan)
TEMPLATE_PAYROLL_LOLOS_GABUNGAN = BASE_DIR / "template" / "template_payroll_lolos.docx"
TEMPLATE_PAYROLL_LOLOS_INDIVIDU = BASE_DIR / "template" / "template_payroll_lolos_individu.docx"
TEMPLATE_PAYROLL_GAGAL_GABUNGAN = BASE_DIR / "template" / "template_payroll_gagal_gabungan.docx"
TEMPLATE_PAYROLL_GAGAL_INDIVIDU = BASE_DIR / "template" / "template_payroll_gagal.docx"

TEMPLATE_CTP_LOLOS_GABUNGAN = BASE_DIR / "template" / "template_ctp_lolos.docx"
TEMPLATE_CTP_LOLOS_INDIVIDU = BASE_DIR / "template" / "template_ctp_lolos_individu.docx"
TEMPLATE_CTP_GAGAL_SCORING = BASE_DIR / "template" / "template_ctp_gagal.docx"
TEMPLATE_CTP_GAGAL_KUOTA = BASE_DIR / "template" / "template_ctp.docx"

OUTPUT_DIR = BASE_DIR / "output"
ZIP_BASE_PATH = BASE_DIR / "Hasil_Surat_Automasi"
ZIP_FILE_PATH = BASE_DIR / "Hasil_Surat_Automasi.zip"
SAMPLE_DATA_PATH = BASE_DIR / "data" / "dummy_data_pengajuan.xlsx"

# ============================================================
# FUNGSI HELPER & FORMAT TEKS OTOMATIS
# ============================================================
def smart_title(text):
    """
    Mengubah teks menjadi Title Case dengan penanganan khusus angka Romawi (I-XX)
    serta singkatan operasional perbankan (KCP, KC, RO, BNM).
    """
    if not text:
        return ""
    text = str(text).title()
    romans = [
        "I", "Ii", "Iii", "Iv", "V", "Vi", "Vii", "Viii", "Ix", "X",
        "Xi", "Xii", "Xiii", "Xiv", "Xv", "Xvi", "Xvii", "Xviii", "Xix", "Xx"
    ]
    for r in romans:
        text = re.sub(rf'\b{r}\b', r.upper(), text)
    text = text.replace("Kcp ", "KCP ").replace("Kc ", "KC ").replace("Ro ", "RO ").replace("Bnm ", "BNM ")
    return text

def clean_text(value):
    """Membersihkan nilai string atau angka tanpa desimal berlebih."""
    if pd.isna(value):
        return ""
    if hasattr(value, "is_integer") and value.is_integer():
        return str(int(value))
    return str(value).strip()

def normalize(value):
    """Menghilangkan spasi ganda dan mengonversi teks ke huruf kapital."""
    value = clean_text(value).upper()
    value = re.sub(r"\s+", " ", value)
    return value

def parse_tanggal_indo(teks_tanggal):
    """Mengurai string tanggal Indonesia (misal: '15 Oktober 2026') menjadi datetime."""
    if pd.isna(teks_tanggal) or str(teks_tanggal).strip() == "":
        return pd.NaT
    teks = str(teks_tanggal).strip().upper()
    bulan_indo = {
        'JANUARI': '01', 'FEBRUARI': '02', 'MARET': '03', 'APRIL': '04',
        'MEI': '05', 'JUNI': '06', 'JULI': '07', 'AGUSTUS': '08',
        'SEPTEMBER': '09', 'OKTOBER': '10', 'NOVEMBER': '11', 'DESEMBER': '12'
    }
    for b_indo, b_angka in bulan_indo.items():
        if b_indo in teks:
            teks = teks.replace(b_indo, b_angka)
            break
    try:
        return pd.to_datetime(teks, format="%d %m %Y")
    except Exception:
        return pd.to_datetime(teks, errors="coerce")

def format_date(value):
    """Mengubah objek tanggal menjadi format resmi bahasa Indonesia (contoh: 15 Oktober 2026)."""
    if pd.isna(value) or str(value).strip() == "":
        return ""
    dt = pd.to_datetime(value, errors="coerce")
    if pd.isna(dt):
        dt = parse_tanggal_indo(value)
    if pd.isna(dt):
        return clean_text(value).title()
    months = [
        "Januari", "Februari", "Maret", "April", "Mei", "Juni",
        "Juli", "Agustus", "September", "Oktober", "November", "Desember"
    ]
    return f"{dt.day} {months[dt.month - 1]} {dt.year}"

def format_rupiah(value):
    """Mengubah angka numerik menjadi format mata uang Rupiah standar (contoh: Rp150.000.000)."""
    if pd.isna(value) or str(value).strip() == "":
        return ""
    if isinstance(value, (int, float)):
        return f"Rp{int(value):,}".replace(",", ".")
    text = str(value).strip()
    try:
        return f"Rp{int(float(text)):,}".replace(",", ".")
    except ValueError:
        pass
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.rsplit(",", 1)[0]
        else:
            text = text.rsplit(".", 1)[0]
    elif "," in text:
        if len(text.rsplit(",", 1)[1]) < 3:
            text = text.rsplit(",", 1)[0]
    elif "." in text:
        if len(text.rsplit(".", 1)[1]) < 3:
            text = text.rsplit(".", 1)[0]
    digits = re.sub(r"[^\d]", "", text)
    if digits:
        return f"Rp{int(digits):,}".replace(",", ".")
    return text

def get_angka_murni(value):
    """Mengekstrak nilai bilangan bulat murni dari teks angka atau format nominal."""
    if pd.isna(value) or str(value).strip() == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).strip()
    try:
        return int(float(text))
    except ValueError:
        pass
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.rsplit(",", 1)[0]
        else:
            text = text.rsplit(".", 1)[0]
    elif "," in text:
        if len(text.rsplit(",", 1)[1]) < 3:
            text = text.rsplit(",", 1)[0]
    elif "." in text:
        if len(text.rsplit(".", 1)[1]) < 3:
            text = text.rsplit(".", 1)[0]
    digits = re.sub(r"[^\d]", "", text)
    if digits:
        return int(digits)
    return 0

def build_region_lookup(master_df):
    """Membangun kamus pencarian cepat nama kantor wilayah dan alamat dari sheet master."""
    lookup = {}
    for _, item in master_df.iterrows():
        singkatan = normalize(item.get("SINGKATAN", ""))
        nama_lengkap = clean_text(item.get("REGION", ""))
        nama_jalan = clean_text(item.get("NAMA JALAN", ""))
        if singkatan not in lookup:
            lookup[singkatan] = {
                "nama_lengkap": nama_lengkap,
                "alamat": nama_jalan
            }
    return lookup

def format_region_italic(text, is_upper=False):
    """Memformat nama Region Office dengan cetak miring pada kata awalan."""
    if pd.isna(text) or str(text).strip() == "":
        return RichText("")
    val = str(text).strip()
    rt = RichText()
    if "-" in val:
        kiri, kanan = val.split("-", 1)
        kiri_parts = kiri.strip().rsplit(" ", 1)
        teks_kiri_awal = kiri_parts[0] if len(kiri_parts) > 1 else kiri.strip()
        teks_romawi = " " + kiri_parts[1] if len(kiri_parts) > 1 else ""
        if is_upper:
            rt.add(teks_kiri_awal.upper(), italic=True)
            rt.add(teks_romawi.upper() + " ")
            rt.add("- " + kanan.strip().upper())
        else:
            rt.add(smart_title(teks_kiri_awal), italic=True)
            rt.add(teks_romawi.upper() + " ")
            rt.add("- " + smart_title(kanan.strip()))
    else:
        if is_upper:
            rt.add(val.upper(), italic=True)
        else:
            rt.add(smart_title(val), italic=True)
    return rt

def format_region_bold(text, is_upper=True):
    """Memformat nama Region Office dengan cetak tebal dan miring."""
    if pd.isna(text) or str(text).strip() == "":
        return RichText("")
    val = str(text).strip()
    rt = RichText()
    kiri = val.split("-", 1)[0].strip() if "-" in val else val
    kiri_parts = kiri.rsplit(" ", 1)
    teks_kiri_awal = kiri_parts[0] if len(kiri_parts) > 1 else kiri
    teks_romawi = " " + kiri_parts[1] if len(kiri_parts) > 1 else ""
    if is_upper:
        rt.add(teks_kiri_awal.upper(), bold=True, italic=True)
        rt.add(teks_romawi.upper(), bold=True)
    else:
        rt.add(smart_title(teks_kiri_awal), bold=True, italic=True)
        rt.add(teks_romawi.upper(), bold=True)
    return rt

def ambil_nilai_kolom(row, daftar_kolom):
    """Mengambil nilai pertama yang ditemukan dari daftar kemungkinan nama kolom."""
    for col in daftar_kolom:
        if col in row and pd.notna(row[col]) and str(row[col]).strip() != "":
            return row[col]
    return ""

def create_base_mapping(row, data_region_master, singkatan_region, nama_bulan_memo):
    """
    Menyusun konteks dictionary Jinja untuk DocxTemplate.
    Menyediakan tag baru yang bersih sekaligus tag legacy untuk kompatibilitas penuh.
    """
    memo_val = ambil_nilai_kolom(row, ["MEMO KOMITE RISIKO", "MEMO RISIKO", "MEMO CRG"])
    tgl_memo_val = ambil_nilai_kolom(row, ["TANGGAL MEMO RISIKO", "TANGGAL MEMO CRG"])
    no_surat_val = ambil_nilai_kolom(row, ["NO SURAT PERSETUJUAN", "NO SURAT KEPUTUSAN", "NO SURAT CF2"])
    
    no_meka_val = ambil_nilai_kolom(row, ["NO SURAT MEKANISME", "NO SURAT MEKA"])
    tgl_meka_val = ambil_nilai_kolom(row, ["TGL SURAT MEKANISME", "TGL SURAT MEKA"])
    
    return {
        # Tag Bersih Baru
        "MEMO_RISIKO": clean_text(memo_val).title(),
        "TGL_MEMO_RISIKO": format_date(tgl_memo_val),
        "NO_SURAT_KEPUTUSAN": clean_text(no_surat_val),
        "NO_SURAT_PERSETUJUAN": clean_text(no_surat_val),
        
        # Tag Kompatibilitas Legacy
        "MEMO_CRG": clean_text(memo_val).title(),
        "TGL_MEMO_CRG": format_date(tgl_memo_val),
        "NO_SURAT_CF2": clean_text(no_surat_val),
        
        # Informasi Umum & Nasabah
        "BULAN": nama_bulan_memo,
        "REGION_SINGKAT": singkatan_region.upper(),
        "REGION": format_region_italic(data_region_master["nama_lengkap"], is_upper=False),
        "NAMA_JALAN": data_region_master["alamat"].title(),
        "REG": format_region_italic(row.get("REGION"), is_upper=False),
        "REG_KAPITAL": format_region_bold(data_region_master["nama_lengkap"], is_upper=True),
        "CAB": smart_title(row.get("NAMA CABANG")),
        "CAB_K": str(row.get("NAMA CABANG", "")).strip().upper(),
        "NAMA": smart_title(row.get("NAMA NASABAH")),
        "NAMA_K": str(row.get("NAMA NASABAH", "")).strip().upper(),
        "WISE": clean_text(row.get("NO WISE")),
        "SEG": clean_text(row.get("SEGMEN")).upper(),
        "INST": clean_text(row.get("INSTITUSI")).upper(),
        "PLAF": format_rupiah(row.get("PLAFOND")),
        "AKS": format_rupiah(row.get("AKSEPTASI PENDAPATAN")),
        "KETERANGAN_PAYROLL": clean_text(row.get("KETERANGAN PAYROLL")).upper(),
        "KETERANGAN_REKOMENDASI": clean_text(row.get("HASIL SCORING WISE")).title(),
        "NO_SURAT_CABANG": clean_text(row.get("NO SURAT CABANG")),
        "TANGGAL_PERSETUJUAN": format_date(row.get("TANGGAL SURAT CABANG")).title(),
        "NAMA_KAPITAL": clean_text(row.get("NAMA NASABAH")).upper(),
        "CABANG_KAPITAL": clean_text(row.get("NAMA CABANG")).upper(),
        "NO_SURAT_MEKA": "-" if not str(no_meka_val).strip() else str(no_meka_val).strip(),
        "TGL_SURAT_MEKA": "-" if not str(tgl_meka_val).strip() else format_date(tgl_meka_val),
    }

def hitung_status_rekomendasi_baru(row):
    """
    Mengevaluasi keputusan persetujuan berdasarkan status scoring dan jenis payroll.
    """
    ket_rek = str(row.get('KETERANGAN REKOMENDASI', '')).strip().lower()
    ket_payroll = str(row.get('KETERANGAN PAYROLL', '')).strip().lower()
    
    if ket_rek == "" or ket_rek == "nan":
        return "BELUM ADA DOKUMEN"
        
    is_ctp = "ctp" in ket_payroll or "commit" in ket_payroll
    is_payroll = "payroll" in ket_payroll or "payrol" in ket_payroll
    
    if is_ctp:
        if "not recommended" in ket_rek:
            return "TIDAK REKOMENDASI CTP"
        elif "recommended" in ket_rek:
            return "REKOMENDASI CTP"
        else:
            return "TIDAK REKOMENDASI LAINNYA"
    elif is_payroll:
        if "not recommended" in ket_rek:
            return "TIDAK REKOMENDASI PAYROLL"
        elif "recommended" in ket_rek:
            return "REKOMENDASI PAYROLL"
        else:
            return "TIDAK REKOMENDASI LAINNYA"
    else:
        return "TIDAK REKOMENDASI LAINNYA"

# ============================================================
# ANTARMUKA WEB STREAMLIT
# ============================================================
def main():
    st.set_page_config(
        page_title="AutoDoc - Automasi Surat & Monitoring Kuota",
        page_icon="📝",
        layout="centered"
    )
    
    # Bilah Sisi (Sidebar)
    with st.sidebar:
        st.header("📌 Tentang Aplikasi")
        st.markdown(
            """
            **AutoDoc Engine** adalah platform automasi dokumen resmi pembiayaan dan monitoring kuota berbasis web.
            
            **Fitur Utama:**
            - 📑 **Pencetakan Batch**: Menghasilkan dokumen Word (`.docx`) siap cetak secara massal.
            - 📊 **Mesin Kuota Real-Time**: Validasi kuota bulanan mencegah alokasi berlebih.
            - ⚡ **Tipografi Otomatis**: Format tanggal Indonesia, mata uang Rupiah, dan penomoran Romawi.
            - 📦 **Ekspor Sekali Klik**: Mengemas seluruh dokumen langsung ke dalam arsip ZIP.
            """
        )
        st.divider()
        st.subheader("📥 Data Uji Coba")
        if SAMPLE_DATA_PATH.exists():
            with open(SAMPLE_DATA_PATH, "rb") as sample_fp:
                st.download_button(
                    label="Unduh Contoh Excel (.xlsx)",
                    data=sample_fp,
                    file_name="dummy_data_pengajuan.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    help="Klik untuk mengunduh dataset contoh buatan untuk pengujian aplikasi."
                )
        st.caption("🔒 Seluruh data contoh dan template bersifat 100% data sintetis/anonim untuk keperluan portofolio.")

    st.title("📝 AutoDoc Engine")
    st.subheader("Sistem Pembuat Dokumen Resmi & Monitoring Kuota Pembiayaan")
    st.markdown("Sistem ini memproses data pengajuan nasabah secara batch, mengelompokkan surat berdasarkan wilayah, dan menghitung kuota bulanan secara akumulatif.")
    st.markdown("---")

    file_pengajuan = st.file_uploader("Unggah Berkas Excel Utama (.xlsx):", type=["xlsx"])
    
    if file_pengajuan:
        excel_file = pd.ExcelFile(file_pengajuan)
        st.markdown("#### 1. Konfigurasi Sheet Excel")
        col_s1, col_s2, col_s3 = st.columns(3)
        with col_s1:
            pilihan_sheet = st.selectbox("Sheet Data Nasabah", excel_file.sheet_names, index=0)
        with col_s2:
            pilihan_sheet_alamat = st.selectbox("Sheet Master Alamat", excel_file.sheet_names, index=1 if len(excel_file.sheet_names) > 1 else 0)
        with col_s3:
            pilihan_sheet_kuota = st.selectbox("Sheet Master Kuota", excel_file.sheet_names, index=2 if len(excel_file.sheet_names) > 2 else 0)

        st.markdown("#### 2. Periode Pemrosesan & Rentang Tanggal")
        col_p1, col_p2 = st.columns(2)
        with col_p1:
            daftar_bulan = [
                "Januari", "Februari", "Maret", "April", "Mei", "Juni",
                "Juli", "Agustus", "September", "Oktober", "November", "Desember"
            ]
            current_month_index = datetime.now().month - 1
            bulan_terpilih = st.selectbox("Bulan Pengajuan", daftar_bulan, index=current_month_index)
        with col_p2:
            rentang_tanggal = st.date_input("Rentang Tanggal Cetak (Mulai & Akhir)", value=[])
        st.markdown("---")

        st.markdown("#### 3. 📋 Seleksi Data Nasabah (Pilih yang Akan Dicetak)")
        try:
            _preview = pd.read_excel(file_pengajuan, sheet_name=pilihan_sheet)
            _preview.columns = [str(c).strip().upper() for c in _preview.columns]
            _preview = _preview.dropna(subset=['NAMA NASABAH'])
            
            # Deteksi kolom tanggal surat (mendukung nama baru dan legacy)
            col_tgl_nama = "TGL SURAT PERSETUJUAN" if "TGL SURAT PERSETUJUAN" in _preview.columns else (
                "TGL SURAT KEPUTUSAN" if "TGL SURAT KEPUTUSAN" in _preview.columns else "TGL SURAT CF2"
            )
            
            _preview['TANGGAL_SISTEM'] = _preview[col_tgl_nama].apply(parse_tanggal_indo)
            _preview = _preview.dropna(subset=['TANGGAL_SISTEM'])
            
            # Perbarui status rekomendasi otomatis dengan mesin keputusan
            _preview['STATUS REKOMENDASI'] = _preview.apply(hitung_status_rekomendasi_baru, axis=1)
            
            if len(rentang_tanggal) == 2:
                _tm = pd.to_datetime(rentang_tanggal[0])
                _ta = pd.to_datetime(rentang_tanggal[1])
                _pengajuan_ui = _preview[(_preview['TANGGAL_SISTEM'] >= _tm) & (_preview['TANGGAL_SISTEM'] <= _ta)].copy()
            else:
                _pengajuan_ui = _preview.copy()
            
            _pengajuan_ui = _pengajuan_ui.reset_index(drop=True)
            
            mode_cetak = st.radio(
                "Pilih Mode Cetak:",
                ["Cetak Semua Data (Otomatis)", "Pilih Nasabah Tertentu (Centang Manual)"],
                horizontal=True
            )
            
            cols_to_show = ['PILIH_CETAK', col_tgl_nama, 'NAMA NASABAH', 'NAMA CABANG', 'REGION', 'STATUS REKOMENDASI']
            
            if mode_cetak == "Pilih Nasabah Tertentu (Centang Manual)":
                _pengajuan_ui.insert(0, 'PILIH_CETAK', False)
                edited_df = st.data_editor(
                    _pengajuan_ui[cols_to_show],
                    use_container_width=True,
                    hide_index=True,
                    column_config={col_tgl_nama: st.column_config.Column("Tgl Surat")},
                    disabled=[col_tgl_nama, 'NAMA NASABAH', 'NAMA CABANG', 'REGION', 'STATUS REKOMENDASI']
                )
            else:
                _pengajuan_ui.insert(0, 'PILIH_CETAK', True)
                edited_df = _pengajuan_ui.copy()
                st.dataframe(
                    _pengajuan_ui[cols_to_show],
                    use_container_width=True,
                    hide_index=True,
                    column_config={col_tgl_nama: st.column_config.Column("Tgl Surat")}
                )
            
            jumlah_cetak = edited_df['PILIH_CETAK'].sum()
            st.info(f"🖨️ Total nasabah yang dipilih untuk dicetak saat ini: **{jumlah_cetak}** dokumen")
            
        except Exception as err:
            edited_df = pd.DataFrame()
            st.info(f"Menunggu berkas Excel yang valid... ({err})")

        st.markdown("---")

    if st.button("🚀 Proses Dokumen & Hitung Kuota", use_container_width=True, type="primary"):
        if not file_pengajuan:
            st.error("Silakan unggah berkas Excel terlebih dahulu!")
            return
            
        if 'edited_df' not in locals() or edited_df.empty or not edited_df['PILIH_CETAK'].any():
            st.error("Silakan pilih/centang minimal 1 nasabah pada tabel di atas!")
            return

        with st.spinner("Membaca data dan mempersiapkan mesin automasi..."):
            try:
                # 1. BACA DATA PENGAJUAN LENGKAP
                pengajuan_full = pd.read_excel(file_pengajuan, sheet_name=pilihan_sheet)
                pengajuan_full.columns = [str(c).strip().upper() for c in pengajuan_full.columns]
                pengajuan_full = pengajuan_full.dropna(subset=['NAMA NASABAH'])
                
                col_tgl_nama = "TGL SURAT PERSETUJUAN" if "TGL SURAT PERSETUJUAN" in pengajuan_full.columns else (
                    "TGL SURAT KEPUTUSAN" if "TGL SURAT KEPUTUSAN" in pengajuan_full.columns else "TGL SURAT CF2"
                )
                
                pengajuan_full['TANGGAL_SISTEM'] = pengajuan_full[col_tgl_nama].apply(parse_tanggal_indo)
                pengajuan_full = pengajuan_full.dropna(subset=['TANGGAL_SISTEM'])
                pengajuan_full['STATUS REKOMENDASI'] = pengajuan_full.apply(hitung_status_rekomendasi_baru, axis=1)
                
                # 2. BACA DATA KUOTA & ALAMAT
                df_kuota = pd.read_excel(file_pengajuan, sheet_name=pilihan_sheet_kuota)
                df_kuota.columns = [str(c).strip().upper() for c in df_kuota.columns]
                
                # 3. AMBIL KUOTA AWAL BULAN
                kuota_awal_non_ctp = 0
                kuota_awal_ctp = 0
                memo_risiko = "-"
                
                baris_bulan = df_kuota[df_kuota["BULAN"].astype(str).str.strip().str.upper() == bulan_terpilih.upper()]
                if not baris_bulan.empty:
                    col_k_payroll = "KUOTA LSO NON CTP" if "KUOTA LSO NON CTP" in baris_bulan.columns else "KUOTA PAYROLL"
                    col_k_ctp = "KUOTA LSO CTP" if "KUOTA LSO CTP" in baris_bulan.columns else "KUOTA CTP"
                    col_memo = "MEMO KOMITE RISIKO" if "MEMO KOMITE RISIKO" in baris_bulan.columns else (
                        "MEMO RISIKO" if "MEMO RISIKO" in baris_bulan.columns else "MEMO CRG"
                    )
                    
                    kuota_awal_non_ctp = int(get_angka_murni(baris_bulan.iloc[0].get(col_k_payroll, 0)))
                    kuota_awal_ctp = int(get_angka_murni(baris_bulan.iloc[0].get(col_k_ctp, 0)))
                    memo_risiko = clean_text(baris_bulan.iloc[0].get(col_memo, "-"))

                # 4. LOGIKA AKUMULASI PENGGUNAAN HISTORIS
                terpakai_lalu_ctp = 0
                terpakai_lalu_non_ctp = 0
                
                if len(rentang_tanggal) == 2:
                    tgl_mulai = pd.to_datetime(rentang_tanggal[0])
                    tgl_akhir = pd.to_datetime(rentang_tanggal[1])
                    
                    idx_bulan_terpilih = daftar_bulan.index(bulan_terpilih) + 1
                    if tgl_mulai.month != idx_bulan_terpilih:
                        nama_bulan_tgl = daftar_bulan[tgl_mulai.month - 1] if 1 <= tgl_mulai.month <= 12 else tgl_mulai.month
                        st.warning(f"⚠️ Perhatian: Anda menarik kuota induk bulan **{bulan_terpilih}**, namun rentang tanggal dimulai pada bulan **{nama_bulan_tgl}**.")

                    df_masa_lalu = pengajuan_full[
                        (pengajuan_full['TANGGAL_SISTEM'].dt.month == tgl_mulai.month) & 
                        (pengajuan_full['TANGGAL_SISTEM'].dt.year == tgl_mulai.year) & 
                        (pengajuan_full['TANGGAL_SISTEM'] < tgl_mulai)
                    ]
                    status_masa_lalu = df_masa_lalu['STATUS REKOMENDASI'].astype(str).str.strip().str.upper()
                    terpakai_lalu_non_ctp = len(df_masa_lalu[status_masa_lalu == 'REKOMENDASI PAYROLL'])
                    terpakai_lalu_ctp = len(df_masa_lalu[status_masa_lalu == 'REKOMENDASI CTP'])

                    pengajuan = pengajuan_full[
                        (pengajuan_full['TANGGAL_SISTEM'] >= tgl_mulai) & 
                        (pengajuan_full['TANGGAL_SISTEM'] <= tgl_akhir)
                    ]
                else:
                    pengajuan = pengajuan_full

                pengajuan = pengajuan.reset_index(drop=True)
                pengajuan['PILIH_CETAK'] = edited_df['PILIH_CETAK'].values

                if OUTPUT_DIR.exists():
                    shutil.rmtree(OUTPUT_DIR)
                OUTPUT_DIR.mkdir(exist_ok=True)
                if ZIP_FILE_PATH.exists():
                    os.remove(ZIP_FILE_PATH)

                pengajuan = pengajuan.sort_values(by='TANGGAL_SISTEM', ascending=True).reset_index(drop=True)
                berhasil, gagal = 0, 0
                terpakai_hari_ini_non_ctp = 0
                terpakai_hari_ini_ctp = 0 

            except Exception as e:
                st.error(f"Gagal mempersiapkan sistem automasi: {e}")
                return

        progress_bar = st.progress(0)
        status_text = st.empty()
        
        total_data_aktual = len(pengajuan) if len(pengajuan) > 0 else 1 
        def update_progress(b, g):
            proses = b + g
            pct = int((proses / total_data_aktual) * 100)
            if pct > 100: pct = 100
            progress_bar.progress(pct)
            status_text.text(f"Memproses {proses} dari {len(pengajuan)} dokumen...")

        df_region = pd.read_excel(file_pengajuan, sheet_name=pilihan_sheet_alamat)
        df_region.columns = [str(c).strip().upper() for c in df_region.columns]
        region_lookup = build_region_lookup(df_region)

        # ============================================================
        # 5. LOGIKA PEMBAGIAN JALUR KUOTA (TIDAK ADA KUOTA MINUS)
        # ============================================================
        pengajuan['TGL_SURAT_BERSIH'] = pengajuan['TANGGAL_SISTEM'].apply(lambda x: format_date(x))
        pengajuan['STATUS_FINAL_CETAK'] = ""
        
        for idx, row in pengajuan.iterrows():
            status_excel = str(row.get('STATUS REKOMENDASI', '')).strip().upper()

            # --- JALUR PAYROLL ---
            if status_excel == "REKOMENDASI PAYROLL":
                sisa_kuota_non_ctp = kuota_awal_non_ctp - (terpakai_lalu_non_ctp + terpakai_hari_ini_non_ctp)
                if sisa_kuota_non_ctp > 0:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_PAYROLL_LOLOS"
                    terpakai_hari_ini_non_ctp += 1
                else:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_PAYROLL_GAGAL"
                
            elif status_excel == "TIDAK REKOMENDASI PAYROLL":
                pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_PAYROLL_GAGAL"

            # --- JALUR CTP (COMMITTED TO PAYROLL) ---
            elif status_excel == "REKOMENDASI CTP":
                sisa_kuota_ctp = kuota_awal_ctp - (terpakai_lalu_ctp + terpakai_hari_ini_ctp)
                if sisa_kuota_ctp > 0:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_CTP_LOLOS"
                    terpakai_hari_ini_ctp += 1
                else:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_CTP_GAGAL_KUOTA"
                    
            elif status_excel == "TIDAK REKOMENDASI CTP":
                sisa_kuota_ctp = kuota_awal_ctp - (terpakai_lalu_ctp + terpakai_hari_ini_ctp)
                if sisa_kuota_ctp > 0:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_CTP_GAGAL_SCORING"
                else:
                    pengajuan.at[idx, 'STATUS_FINAL_CETAK'] = "CETAK_CTP_GAGAL_KUOTA"

        pengajuan_cetak = pengajuan[pengajuan['PILIH_CETAK'] == True]

        if pengajuan_cetak.empty:
            st.error("Silakan centang minimal 1 nasabah pada tabel!")
            return

        # DOKUMEN GABUNGAN (Payroll Lolos, CTP Lolos, dan Payroll Gagal)
        data_lolos = pengajuan_cetak[pengajuan_cetak['STATUS_FINAL_CETAK'].isin(["CETAK_PAYROLL_LOLOS", "CETAK_CTP_LOLOS", "CETAK_PAYROLL_GAGAL"])]

        # DOKUMEN INDIVIDU (CTP Gagal)
        data_gagal = pengajuan_cetak[pengajuan_cetak['STATUS_FINAL_CETAK'].isin(["CETAK_CTP_GAGAL_SCORING", "CETAK_CTP_GAGAL_KUOTA"])]

        urutan_file = 1
        
        # ---------------------------------------------------------
        # MESIN CETAK A: DOKUMEN GABUNGAN & INDIVIDU LOLOS
        # ---------------------------------------------------------
        grup_lolos = data_lolos.groupby(['STATUS_FINAL_CETAK', 'REGION', 'TGL_SURAT_BERSIH'])
        
        for (kategori_cetak, nama_region_excel, tgl_surat_excel), grup_nasabah in grup_lolos:
            try:
                if kategori_cetak == "CETAK_PAYROLL_LOLOS":
                    TPL_GABUNGAN = TEMPLATE_PAYROLL_LOLOS_GABUNGAN
                    TPL_INDIVIDU = TEMPLATE_PAYROLL_LOLOS_INDIVIDU
                    prefix_file = "REKOM_PAYROLL"
                    ket_rekomendasi = "REKOMENDASI"
                elif kategori_cetak == "CETAK_PAYROLL_GAGAL":
                    TPL_GABUNGAN = TEMPLATE_PAYROLL_GAGAL_GABUNGAN
                    TPL_INDIVIDU = TEMPLATE_PAYROLL_GAGAL_INDIVIDU
                    prefix_file = "TOLAK_PAYROLL"
                    ket_rekomendasi = "TIDAK REKOMENDASI"
                else:
                    TPL_GABUNGAN = TEMPLATE_CTP_LOLOS_GABUNGAN
                    TPL_INDIVIDU = TEMPLATE_CTP_LOLOS_INDIVIDU
                    prefix_file = "REKOM_CTP"
                    ket_rekomendasi = "REKOMENDASI"

                row_pertama = grup_nasabah.iloc[0]
                singkatan_region = normalize(nama_region_excel)
                data_region_master = region_lookup.get(singkatan_region, {"nama_lengkap": "-", "alamat": "-"})
                
                tgl_memo_raw = ambil_nilai_kolom(row_pertama, ["TANGGAL MEMO RISIKO", "TANGGAL MEMO CRG"])
                tgl_memo_dt = parse_tanggal_indo(tgl_memo_raw)
                nama_bulan_memo = "-"
                if pd.notna(tgl_memo_dt):
                    nama_bulan_memo = daftar_bulan[tgl_memo_dt.month - 1]

                jumlah_nasabah = len(grup_nasabah)
                
                # JIKA GABUNGAN (>1 NASABAH DALAM SATU REGION & TANGGAL)
                if jumlah_nasabah > 1:
                    daftar_nasabah = []
                    daftar_area_unik = []
                    
                    for _, row in grup_nasabah.iterrows():
                        area_bersih = clean_text(row.get("AREA"))
                        if area_bersih and area_bersih.upper() not in [a.upper() for a in daftar_area_unik]:
                            daftar_area_unik.append(area_bersih)
                            
                        no_meka_item = ambil_nilai_kolom(row, ["NO SURAT MEKANISME", "NO SURAT MEKA"])
                        tgl_meka_item = ambil_nilai_kolom(row, ["TGL SURAT MEKANISME", "TGL SURAT MEKA"])
                        
                        item = {
                            "REG": format_region_italic(row.get("REGION"), is_upper=False),
                            "REG_PLAIN": smart_title(row.get("REGION")) if pd.notna(row.get("REGION")) else "-",
                            "AREA": clean_text(row.get("AREA")).title(),
                            "AREA_PERTAMA_K": str(row.get("AREA", "")).strip().upper(),
                            "CAB": smart_title(row.get("NAMA CABANG")),
                            "CAB_K": str(row.get("NAMA CABANG", "")).strip().upper(),
                            "NO_SURAT_CABANG": clean_text(row.get("NO SURAT CABANG")),
                            "TANGGAL_PERSETUJUAN": format_date(row.get("TANGGAL SURAT CABANG")).title(),
                            "KETERANGAN_PAYROLL": clean_text(row.get("KETERANGAN PAYROLL")).upper(),
                            "KETERANGAN_REKOMENDASI": clean_text(row.get("HASIL SCORING WISE")).title(),
                            "NAMA": smart_title(row.get("NAMA NASABAH")),
                            "NAMA_K": str(row.get("NAMA NASABAH", "")).strip().upper(),
                            "WISE": clean_text(row.get("NO WISE")),
                            "SEG": clean_text(row.get("SEGMEN")).upper(),
                            "INST": clean_text(row.get("INSTITUSI")).upper(),
                            "PLAF": format_rupiah(row.get("PLAFOND")),
                            "AKS": format_rupiah(row.get("AKSEPTASI PENDAPATAN")),
                            "NO_SURAT_MEKA": "-" if not str(no_meka_item).strip() else str(no_meka_item).strip(),
                            "TGL_SURAT_MEKA": "-" if not str(tgl_meka_item).strip() else format_date(tgl_meka_item),
                        }
                        daftar_nasabah.append(item)
                    
                    area_gabungan_title = "\n".join([a.title() for a in daftar_area_unik])
                    area_gabungan_kapital = "\n".join([a.upper() for a in daftar_area_unik])
                    
                    daftar_nama_cabang = [f"{n['NAMA']} (Nasabah {n['CAB']})" for n in daftar_nasabah]
                    def format_list_indo(lst):
                        if not lst: return "-"
                        if len(lst) == 1: return lst[0]
                        if len(lst) == 2: return f"{lst[0]} dan {lst[1]}"
                        return ", ".join(lst[:-1]) + f", dan {lst[-1]}"
                    teks_nama_cabang_gabungan = format_list_indo(daftar_nama_cabang)
                    
                    mapping = create_base_mapping(row_pertama, data_region_master, singkatan_region, nama_bulan_memo)
                    mapping.update({
                        "GABUNGAN_NAMA_CABANG": teks_nama_cabang_gabungan,
                        "daftar_nasabah": daftar_nasabah,         
                        "total_nasabah": len(daftar_nasabah),     
                        "TGL_SURAT_KEPUTUSAN": str(tgl_surat_excel).title(),
                        "TGL_SURAT_PERSETUJUAN": str(tgl_surat_excel).title(),
                        "TGL_SURAT_CF2": str(tgl_surat_excel).title(),
                        "AREA": area_gabungan_title,
                        "AREA_KAPITAL": area_gabungan_kapital,
                        "AREA_PERTAMA": clean_text(row_pertama.get("AREA")).title(),
                        "AREA_PERTAMA_K": str(row_pertama.get("AREA", "")).strip().upper(),
                        "REGION_KAPITAL": format_region_italic(row_pertama.get("REGION"), is_upper=True),
                    })
                    
                    doc = DocxTemplate(TPL_GABUNGAN)
                    
                    # 🔨 PEMBUATAN TABEL SUBDOKUMEN DINAMIS
                    subdoc = doc.new_subdoc()
                    table = subdoc.add_table(rows=1, cols=10)
                    table.style = 'Table Grid'
                    table.autofit = False 
                    table.allow_autofit = False 
                    
                    col_widths = [Cm(1.2), Cm(2.8), Cm(2.8), Cm(3.5), Cm(3.5), Cm(1.5), Cm(3.4), Cm(2.5), Cm(2.5), Cm(1.5)]
                    for col_idx, col in enumerate(table.columns):
                        col.width = col_widths[col_idx] 
                    
                    headers = ["RO", "AREA", "NAMA CABANG", "NAMA NASABAH", "NO WI SE", "SEGMEN", "INSTITUSI", "PLAFOND", "AKSEPTASI PENDAPATAN", "KET"]
                    hdr_cells = table.rows[0].cells
                    for i, hdr_text in enumerate(headers):
                        hdr_cells[i].text = hdr_text
                        hdr_cells[i].width = col_widths[i]
                        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                        
                        # Latar Belakang Header Berwarna Teal (#3DB195)
                        shading_elm = parse_xml(r'<w:shd {} w:fill="3DB195"/>'.format(nsdecls('w')))
                        hdr_cells[i]._tc.get_or_add_tcPr().append(shading_elm)
                        for paragraph in hdr_cells[i].paragraphs:
                            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for run in paragraph.runs:
                                run.font.name = 'Arial'
                                run.font.size = Pt(7.5)
                                run.font.bold = True
                                run.font.color.rgb = RGBColor(255, 255, 255)
                    
                    # Isi Baris Data Nasabah
                    for item_data in daftar_nasabah:
                        row_cells = table.add_row().cells
                        row_values = [
                            item_data.get("REG_PLAIN", ""),
                            item_data.get("AREA_PERTAMA_K", ""),
                            item_data.get("CAB_K", ""),
                            item_data.get("NAMA_K", ""),
                            item_data.get("WISE", ""),
                            item_data.get("SEG", ""),
                            item_data.get("INST", ""),
                            item_data.get("PLAF", ""),
                            item_data.get("AKS", ""),
                            item_data.get("KETERANGAN_PAYROLL", "")
                        ]
                        for i, val in enumerate(row_values):
                            row_cells[i].text = str(val)
                            row_cells[i].width = col_widths[i]
                            row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                            for paragraph in row_cells[i].paragraphs:
                                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                for run in paragraph.runs:
                                    run.font.name = 'Arial'
                                    run.font.size = Pt(7.5)
                    
                    mapping["TABEL_GABUNGAN"] = subdoc
                    doc.render(mapping)
                    nama_file_output = f"{urutan_file}_{prefix_file}_GABUNGAN_{singkatan_region}.docx"
                    
                # JIKA INDIVIDU (1 NASABAH)
                else:
                    mapping = create_base_mapping(row_pertama, data_region_master, singkatan_region, nama_bulan_memo)
                    mapping.update({
                        "TGL_SURAT_KEPUTUSAN": str(tgl_surat_excel).title(),
                        "TGL_SURAT_PERSETUJUAN": str(tgl_surat_excel).title(),
                        "TGL_SURAT_CF2": str(tgl_surat_excel).title(),
                        "AREA": clean_text(row_pertama.get("AREA")).title(),
                        "AREA_KAPITAL": clean_text(row_pertama.get("AREA")).upper(),
                        "AREA_PERTAMA": str(row_pertama.get("AREA", "")).strip().title(),
                        "AREA_PERTAMA_K": str(row_pertama.get("AREA", "")).strip().upper(),
                        "KET": ket_rekomendasi,
                    })
                    doc = DocxTemplate(TPL_INDIVIDU)
                    doc.render(mapping)
                    safe_nama = re.sub(r'[\\/*?:"<>|]', '_', mapping["NAMA"])
                    nama_file_output = f"{urutan_file}_{prefix_file}_INDIVIDU_{safe_nama}.docx"
                
                output_file = OUTPUT_DIR / nama_file_output
                doc.save(output_file)
                urutan_file += 1
                berhasil += jumlah_nasabah 
                update_progress(berhasil, gagal)
                        
            except Exception as e:
                st.error(f"Error saat mencetak Grup {prefix_file} Region {nama_region_excel}: {e}")
                gagal += len(grup_nasabah)
                update_progress(berhasil, gagal)
                
        # ---------------------------------------------------------
        # MESIN CETAK B: NASABAH GAGAL CTP (INDIVIDU)
        # ---------------------------------------------------------
        for index, row in data_gagal.iterrows():
            try:
                status_final_cetak = row.get('STATUS_FINAL_CETAK')
                if status_final_cetak == "CETAK_CTP_GAGAL_KUOTA":
                    template_terpilih = TEMPLATE_CTP_GAGAL_KUOTA
                    label_status = "TIDAK_REKOMENDASI_CTP_KUOTA_HABIS"
                else:
                    template_terpilih = TEMPLATE_CTP_GAGAL_SCORING
                    label_status = "TIDAK_REKOMENDASI_CTP_Scoring_Wise"
                
                tgl_memo_raw = ambil_nilai_kolom(row, ["TANGGAL MEMO RISIKO", "TANGGAL MEMO CRG"])
                tgl_memo_dt = parse_tanggal_indo(tgl_memo_raw)
                nama_bulan_memo = "-"
                if pd.notna(tgl_memo_dt):
                    nama_bulan_memo = daftar_bulan[tgl_memo_dt.month - 1]
                    
                singkatan_region = normalize(row.get("REGION"))
                data_region_master = region_lookup.get(singkatan_region, {"nama_lengkap": "-", "alamat": "-"})
                
                tgl_surat_item = ambil_nilai_kolom(row, ["TGL SURAT PERSETUJUAN", "TGL SURAT KEPUTUSAN", "TGL SURAT CF2"])
                
                mapping = create_base_mapping(row, data_region_master, singkatan_region, nama_bulan_memo)
                mapping.update({
                    "TGL_SURAT_KEPUTUSAN": str(tgl_surat_item).title(),
                    "TGL_SURAT_PERSETUJUAN": str(tgl_surat_item).title(),
                    "TGL_SURAT_CF2": str(tgl_surat_item).title(),
                    "AREA": clean_text(row.get("AREA")).title(),
                    "AREA_KAPITAL": clean_text(row.get("AREA")).upper(),
                    "AREA_PERTAMA": str(row.get("AREA", "")).strip().title(),
                    "AREA_PERTAMA_K": str(row.get("AREA", "")).strip().upper(),
                    "KET": "TIDAK REKOMENDASI",
                })
                
                doc = DocxTemplate(template_terpilih)
                doc.render(mapping)
                
                safe_nama = re.sub(r'[\\/*?:"<>|]', '_', mapping["NAMA"])
                nama_file_output = f"{urutan_file}_{label_status}_{safe_nama}.docx"
                output_file = OUTPUT_DIR / nama_file_output
                doc.save(output_file)
                
                urutan_file += 1
                berhasil += 1
                update_progress(berhasil, gagal)
                
            except Exception as e:
                st.error(f"Error saat mencetak nasabah Gagal {row.get('NAMA NASABAH')}: {e}")
                gagal += 1
                update_progress(berhasil, gagal)
                
        progress_bar.progress(100)
        status_text.text(f"Selesai memproses {berhasil + gagal} dokumen.")

        shutil.make_archive(str(ZIP_BASE_PATH), 'zip', str(OUTPUT_DIR))
        st.success(f"🎉 Pemrosesan Selesai! Berhasil: {berhasil} | Gagal: {gagal}")
        st.markdown("---")

        # ============================================================
        # 6. DASHBOARD METRIK: REKAPITULASI KUOTA FINAL
        # ============================================================
        st.markdown(f"### 📊 Ringkasan Kuota Bulan {bulan_terpilih}")
        
        total_terpakai_bulanan_non_ctp = terpakai_lalu_non_ctp + terpakai_hari_ini_non_ctp
        total_terpakai_bulanan_ctp = terpakai_lalu_ctp + terpakai_hari_ini_ctp

        if total_terpakai_bulanan_non_ctp > kuota_awal_non_ctp:
            total_terpakai_bulanan_non_ctp = kuota_awal_non_ctp
            
        if total_terpakai_bulanan_ctp > kuota_awal_ctp:
            total_terpakai_bulanan_ctp = kuota_awal_ctp

        sisa_akhir_non_ctp = max(0, kuota_awal_non_ctp - total_terpakai_bulanan_non_ctp)
        sisa_akhir_ctp = max(0, kuota_awal_ctp - total_terpakai_bulanan_ctp)

        m1, m2, m3 = st.columns(3)
        m1.metric("Kuota Payroll", kuota_awal_non_ctp)
        m2.metric("Terpakai Payroll", total_terpakai_bulanan_non_ctp)
        m3.metric("Sisa Kuota Payroll", sisa_akhir_non_ctp)

        m4, m5, m6 = st.columns(3)
        m4.metric("Kuota CTP", kuota_awal_ctp)
        m5.metric("Terpakai CTP", total_terpakai_bulanan_ctp)
        m6.metric("Sisa Kuota CTP", sisa_akhir_ctp)

        st.markdown("")

        data_rekap = {
            "BULAN": [bulan_terpilih],
            "MEMO KOMITE RISIKO": [memo_risiko],
            "KUOTA LSO NON CTP": [kuota_awal_non_ctp],
            "KUOTA LSO CTP": [kuota_awal_ctp],
            "PENGGUNA KUOTA LSO NON CTP": [total_terpakai_bulanan_non_ctp],
            "PENGGUNA KUOTA LSO CTP": [total_terpakai_bulanan_ctp],
            "SISA KUOTA LSO NON CTP": [sisa_akhir_non_ctp],
            "SISA KUOTA LSO CTP": [sisa_akhir_ctp]
        }
        df_rekap = pd.DataFrame(data_rekap)
        st.dataframe(df_rekap, use_container_width=True)

        if ZIP_FILE_PATH.exists():
            with open(ZIP_FILE_PATH, "rb") as fp:
                st.download_button(
                    "📦 Unduh Seluruh Surat (ZIP)",
                    data=fp,
                    file_name="Hasil_Surat_Automasi.zip",
                    mime="application/zip",
                    type="primary",
                    use_container_width=True
                )

if __name__ == "__main__":
    main()
